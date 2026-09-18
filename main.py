import io
import os
import re
import sys
import json
import queue
import threading
import asyncio
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi import FastAPI
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from dotenv import load_dotenv

# Load env from project directory
load_dotenv(Path(__file__).parent / ".env")

# Add src to path
sys.path.insert(0, str(Path(__file__).parent))

app = FastAPI(title="Point Blank Operation")

static_dir = Path(__file__).parent / "static"
app.mount("/static", StaticFiles(directory=str(static_dir)), name="static")

# Single global stop flag — only one pipeline runs at a time
_stop_event = threading.Event()


class ResearchRequest(BaseModel):
    niche: str
    state: str
    city: Optional[str] = None
    resume_state: Optional[dict] = None


@app.get("/", response_class=HTMLResponse)
async def home():
    return (Path(__file__).parent / "static" / "index.html").read_text()


@app.post("/api/stop")
async def stop_research():
    _stop_event.set()
    return JSONResponse({"ok": True})


@app.post("/api/run")
async def run_research(request: ResearchRequest):
    _stop_event.clear()
    update_queue: queue.Queue = queue.Queue()

    def run_pipeline():
        try:
            from src.rank_rent.crew import RankRentPipeline, StoppedError

            pipeline = RankRentPipeline(
                niche=request.niche,
                state=request.state,
                city=request.city or None,
                stop_event=_stop_event,
                resume_state=request.resume_state or {},
                update_callback=lambda phase, status, data: update_queue.put(
                    {"phase": phase, "status": status, "data": data}
                ),
            )
            result = pipeline.run()
            if result.get("no_cities"):
                update_queue.put({"phase": "done", "status": "complete", "data": "no_cities"})
            else:
                update_queue.put({"phase": "done", "status": "complete", "data": "All phases complete!"})
        except StoppedError:
            update_queue.put({"phase": "stopped", "status": "stopped", "data": "Research stopped."})
        except Exception as exc:
            update_queue.put({"phase": "error", "status": "error", "data": str(exc)})

    thread = threading.Thread(target=run_pipeline, daemon=True)
    thread.start()

    async def generate():
        loop = asyncio.get_event_loop()
        while True:
            try:
                msg = await loop.run_in_executor(
                    None, lambda: update_queue.get(timeout=20)
                )
                yield f"data: {json.dumps(msg)}\n\n"
                if msg.get("phase") in ("done", "error", "stopped"):
                    break
            except queue.Empty:
                # keepalive ping so connection stays open
                yield f"data: {json.dumps({'phase': 'ping', 'status': 'waiting', 'data': ''})}\n\n"

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )


class ExportRequest(BaseModel):
    text: str
    niche: str = ""
    city: str = ""
    state: str = ""


def _parse_prospects(text: str) -> list[dict]:
    """Parse Phase 5 prospect text into a list of row dicts."""
    prospects = []
    # Split on prospect blocks: lines starting with #N —
    blocks = re.split(r'\n(?=#\d+\s*[—\-])', text)
    for block in blocks:
        block = block.strip()
        if not block or not re.match(r'^#\d+', block):
            continue

        def _field(pattern, default=""):
            m = re.search(pattern, block, re.IGNORECASE)
            return m.group(1).strip() if m else default

        name_m = re.match(r'^#\d+\s*[—\-]\s*(.+)', block)
        name = name_m.group(1).strip() if name_m else ""

        prospects.append({
            "Business Name": name,
            "Website": _field(r'Website:\s*(.+)'),
            "Phone": _field(r'Phone:\s*(.+)'),
            "Found Via": _field(r'Found Via:\s*(.+)'),
            "Owner": _field(r'Owner:\s*(.+)'),
            "Notes": _field(r'Notes:\s*(.+)'),
            "Called?": "",
            "Call Notes": "",
        })
    return prospects


@app.post("/api/export-prospects")
async def export_prospects(req: ExportRequest):
    rows = _parse_prospects(req.text)
    if not rows:
        return JSONResponse({"error": "No prospects found to export"}, status_code=400)

    wb = openpyxl.Workbook()
    ws = wb.active
    label = f"{req.niche.title()} — {req.city}, {req.state}".strip(" —,")
    ws.title = label[:31] or "Prospects"

    # Header row
    headers = ["#", "Business Name", "Website", "Phone", "Found Via", "Owner", "Notes", "Called?", "Call Notes"]
    header_fill = PatternFill("solid", fgColor="1F6FEB")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Data rows with alternating shading
    fill_a = PatternFill("solid", fgColor="0D1117")
    fill_b = PatternFill("solid", fgColor="161B22")
    font_data = Font(color="E6EDF3")

    for i, row in enumerate(rows, 1):
        fill = fill_a if i % 2 == 1 else fill_b
        values = [i, row["Business Name"], row["Website"], row["Phone"],
                  row["Found Via"], row["Owner"], row["Notes"], "", ""]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=i + 1, column=col, value=val)
            cell.fill = fill
            cell.font = font_data
            cell.alignment = Alignment(wrap_text=True)

    # Column widths
    widths = [4, 28, 32, 16, 18, 14, 36, 10, 36]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"prospects_{req.niche.replace(' ', '_')}_{req.city.replace(' ', '_')}_{req.state}.xlsx".lower()
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/health")
async def health():
    keys = {
        "anthropic": bool(os.getenv("ANTHROPIC_API_KEY")),
        "semrush": bool(os.getenv("SEMRUSH_API_KEY")),
        "serper": bool(os.getenv("SERPER_API_KEY")),
    }
    return {"status": "ok", "api_keys_loaded": keys}
